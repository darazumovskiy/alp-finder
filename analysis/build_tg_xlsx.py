#!/usr/bin/env python3
"""Excel-выгрузка активности волонтёров TG-группы.

Те же данные, что в otchet-tg-aktivnost.html (build_tg_report.py):
лист «Авторы» — активность и вклад, лист «Молчуны» — участники без сообщений.

Запуск: analysis/.venv/bin/python analysis/build_tg_xlsx.py
Выход:  analysis/tg-aktivnost.xlsx
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
A = ROOT / "analysis"
OUT = A / "tg-aktivnost.xlsx"

agg = json.loads((A / "tg_activity.json").read_text(encoding="utf-8"))
participants = json.loads((A / "tg_participants.json").read_text(encoding="utf-8"))
scores = {s["author"]: s for s in
          json.loads((A / "tg_scores.json").read_text(encoding="utf-8"))}
silent = [m for m in participants if m["messages"] == 0]

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()

ws = wb.active
ws.title = "Авторы"
sheet(
    ws,
    ["Автор", "Сообщений", "Дней активности", "Первый день", "Последний день",
     "Часы (оценка)", "Сессий", "Фото всего", "Кандидаты (фото)", "Принято",
     "Обсуждения (0-10, LLM)", "Комментарий (LLM)"],
    [[a["author"], a["messages"], a["days_active"], a["first_day"], a["last_day"],
      a["est_hours"], a["sessions"], a["media_total"], a["candidates"], a["accepted"],
      scores[a["author"]]["discussion"] if a["author"] in scores else None,
      scores[a["author"]]["comment"] if a["author"] in scores else None]
     for a in agg],
    [26, 11, 14, 12, 14, 12, 8, 10, 15, 9, 20, 80],
)

ws2 = wb.create_sheet("Молчуны")
sheet(
    ws2,
    ["Имя", "Username", "Вступил(а)"],
    [[m["name"], f"@{m['username']}" if m["username"] else None, m["joined"]]
     for m in sorted(silent, key=lambda m: m["joined"] or "")],
    [30, 24, 18],
)

n_msgs = sum(a["messages"] for a in agg)
ws3 = wb.create_sheet("Сводка")
sheet(
    ws3,
    ["Показатель", "Значение"],
    [["Период", f"{min(a['first_day'] for a in agg)} — {max(a['last_day'] for a in agg)}"],
     ["Участников в группе", len(participants)],
     ["Писали сообщения", len(agg)],
     ["Молчуны (0 сообщений)", len(silent)],
     ["Сообщений в темах группы", n_msgs],
     ["Принятых кадров (кластеров)", sum(a["accepted"] for a in agg)],
     ["Методика", "см. https://alp-finder-team.pages.dev/otchet-tg-aktivnost.html"]],
    [30, 60],
)

wb.save(OUT)
print(f"ok: {OUT} ({OUT.stat().st_size // 1024} КБ)")
